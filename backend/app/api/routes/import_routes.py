import os
import tempfile
import csv
from datetime import time
from pathlib import Path

from openpyxl import load_workbook
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field, model_validator
from sqlalchemy.orm import Session

from app.database.connection import get_db
from app.models.enums import DayOfWeek
from app.models.person import Person
from app.services.import_schedule_service import import_schedule
from app.services.parsers.generic_schedule_parser import parse_dataframe
from app.services.parsers.image_parser import extract_ocr_data_from_image, extract_text_from_image
from app.services.parsers.image_schedule_parser import parse_image_to_schedule
from app.services.parsers.ocr_schedule_parser import parse_schedule_ocr_data
from app.services.parsers.pdf_parser import extract_tables, extract_text
from app.services.validators.schedule_validator import validate_schedule

router = APIRouter(prefix="/import", tags=["Import"])
SUPPORTED_EXTENSIONS = {"csv", "xlsx", "pdf", "png", "jpg", "jpeg"}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024


class UploadRejected(ValueError):
    """Raised when the upload itself is invalid, before parsing begins."""


class EditableScheduleRow(BaseModel):
    activity: str = Field(min_length=1, max_length=100)
    day: DayOfWeek
    start: time
    end: time

    @model_validator(mode="after")
    def valid_range(self):
        if self.end <= self.start:
            raise ValueError("End time must be after start time")
        return self


class ScheduleImportRequest(BaseModel):
    person_id: int
    schedule: list[EditableScheduleRow]
    replace_existing: bool = False


def _read_upload(file: UploadFile) -> bytes:
    file.file.seek(0)
    content = file.file.read(MAX_UPLOAD_BYTES + 1)
    if not content:
        raise UploadRejected("El archivo está vacío")
    if len(content) > MAX_UPLOAD_BYTES:
        raise UploadRejected("El archivo supera el límite de 20 MB")
    return content


def _read_csv(path: str):
    raw = Path(path).read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(text.splitlines(), dialect=dialect))


def _rows_from_tables(tables):
    rows = []
    for table in tables:
        if not table or len(table) < 2:
            continue
        headers = [str(value or "").strip() for value in table[0]]
        rows.extend(
            dict(zip(headers, row))
            for row in table[1:]
            if row and any(value not in (None, "") for value in row)
        )
    return rows


def load_dataframe(file: UploadFile):
    suffix = Path(file.filename or "").suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{suffix}") as temporary:
            temporary.write(_read_upload(file))
            path = temporary.name

        if suffix == "csv":
            return _read_csv(path)
        if suffix == "xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            workbook.close()
            if not values:
                return []
            headers = [str(value or "").strip() for value in values[0]]
            return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
        if suffix == "pdf":
            rows = _rows_from_tables(extract_tables(path))
            if rows:
                return rows
            schedule = parse_image_to_schedule(extract_text(path))
            return schedule
        schedule = parse_schedule_ocr_data(extract_ocr_data_from_image(path))
        return schedule or parse_image_to_schedule(extract_text_from_image(path))
    finally:
        if path and os.path.exists(path):
            os.unlink(path)


def _preview(file: UploadFile):
    try:
        loaded = load_dataframe(file)
        # OCR parsers already produce canonical rows; spreadsheets still need
        # column detection and normalization.
        result = {"success": True, "schedule": loaded} if isinstance(loaded, list) and loaded and set(loaded[0]).issuperset({"activity", "day", "start", "end"}) else parse_dataframe(loaded)
        if not result["success"]:
            raise HTTPException(422, result["error"])
        validation = validate_schedule(result["schedule"])
        return {
            "success": validation["valid"],
            "schedule": validation["cleaned"],
            "errors": validation["errors"],
        }
    except HTTPException:
        raise
    except UploadRejected as error:
        status_code = 413 if "20 MB" in str(error) else 422
        raise HTTPException(status_code, str(error)) from error
    except Exception as error:
        suffix = Path(file.filename or "").suffix.lower()
        if suffix in {".png", ".jpg", ".jpeg", ".pdf"}:
            return {
                "success": False,
                "schedule": [{
                    "activity": "",
                    "day": DayOfWeek.MONDAY,
                    "start": time(8, 0),
                    "end": time(9, 0),
                }],
                "errors": [{
                    "type": "OCR_MANUAL_REVIEW",
                    "message": "No se pudo reconocer automáticamente el documento. Completa o corrige la fila mostrada antes de guardar.",
                }],
            }
        raise HTTPException(422, f"No se pudo interpretar el archivo: {error}") from error


@router.post("/excel")
async def import_excel(file: UploadFile = File(...)):
    return _preview(file)


@router.post("/preview")
async def preview_file(file: UploadFile = File(...)):
    return _preview(file)


@router.post("/save")
def save_edited_schedule(payload: ScheduleImportRequest, db: Session = Depends(get_db)):
    person = db.get(Person, payload.person_id)
    if not person:
        raise HTTPException(404, "Person not found")
    if payload.replace_existing:
        for activity in list(person.activities):
            db.delete(activity)
        db.flush()
    rows = [row.model_dump() for row in payload.schedule]
    return {"success": True, "imported": import_schedule(payload.person_id, rows, db)}


@router.post("/excel/save")
async def import_and_save(person_id: int = Form(...), file: UploadFile = File(...), db: Session = Depends(get_db)):
    preview = _preview(file)
    if not preview["success"]:
        return preview
    if not db.get(Person, person_id):
        raise HTTPException(404, "Person not found")
    return {"success": True, "imported": import_schedule(person_id, preview["schedule"], db)}
